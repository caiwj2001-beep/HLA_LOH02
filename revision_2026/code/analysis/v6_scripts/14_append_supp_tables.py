#!/usr/bin/env python3
"""Append Supplementary Table 5 (panel sensitivity) + Table 6 (mediation, if ready)
to Supplementary_Tables.docx, preserving existing 4 tables. Idempotent rebuild."""
import csv, os
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

TAB="C:/Users/Administrator/科研文件/HLA_LOH02/V6/04_processed_data/results"
OUTd="C:/Users/Administrator/科研文件/HLA_LOH02/V6/05_tables"
NAVY=RGBColor(0x00,0x33,0x66)

def load(f):
    with open(os.path.join(TAB,f),encoding="utf-8") as fh:
        return list(csv.reader(fh,delimiter="\t"))
def shade(cell,hexc):
    tcPr=cell._tc.get_or_add_tcPr(); sh=OxmlElement("w:shd")
    sh.set(qn("w:val"),"clear"); sh.set(qn("w:color"),"auto"); sh.set(qn("w:fill"),hexc); tcPr.append(sh)
def setc(cell,text,bold=False,white=False,size=9,align="left"):
    cell.text=""; p=cell.paragraphs[0]
    p.alignment={"left":WD_ALIGN_PARAGRAPH.LEFT,"center":WD_ALIGN_PARAGRAPH.CENTER}[align]
    r=p.add_run(str(text).strip().strip('"')); r.font.size=Pt(size); r.font.name="Arial"; r.bold=bold
    if white:r.font.color.rgb=RGBColor(0xFF,0xFF,0xFF)
    elif bold:r.font.color.rgb=NAVY
def mktable(doc,rows,title,cap,fs=9):
    h=doc.add_paragraph(); rr=h.add_run(title); rr.bold=True; rr.font.color.rgb=NAVY; rr.font.size=Pt(11); rr.font.name="Arial"
    cp=doc.add_paragraph(); cr=cp.add_run(cap); cr.italic=True; cr.font.size=Pt(8.5); cr.font.name="Arial"
    t=doc.add_table(rows=1,cols=len(rows[0])); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.style="Table Grid"
    for j,c in enumerate(rows[0]):
        setc(t.rows[0].cells[j],c,bold=True,white=True,size=fs,align="center"); shade(t.rows[0].cells[j],"003366")
    for i,row in enumerate(rows[1:]):
        cells=t.add_row().cells
        for j,c in enumerate(row): setc(cells[j],c,size=fs,align="left" if j==0 else "center")
        if i%2==1:
            for c in cells: shade(c,"EAF0F6")
    doc.add_paragraph()

# Rebuild the whole supplementary doc from scratch (KISS: reuse 07 logic via import-free rebuild)
import subprocess
# regenerate the base 4-table doc first
subprocess.run(["python","C:/Users/Administrator/科研文件/HLA_LOH02/V6/03_scripts/07_build_tables.py"],check=True)

doc=Document(os.path.join(OUTd,"Supplementary_Tables.docx"))

# ST5: panel sensitivity
p5=load("SuppTable5_panel_sensitivity.tsv")
hdr=["Effect","Full-cohort OR","Restricted OR","95% CI low","95% CI high","P value"]
rows=[hdr]
for r in p5[1:]:
    eff,full,ror,lo,hi,p=r
    pv="<0.001" if float(p)<0.001 else f"{float(p):.3f}"
    rows.append([eff,full,ror,lo,hi,pv])
mktable(doc,rows,
  "Supplementary Table 5. Panel-restricted sensitivity analysis (IMPACT468/505).",
  "Headline associations re-estimated in the 24,512 tumors (82.3% of the primary cohort) profiled "
  "with IMPACT468 or IMPACT505, excluding earlier panel versions. Effect sizes are essentially "
  "unchanged from the full cohort, indicating robustness to panel/batch differences. OR, odds ratio.",
  fs=8)

# ST6: mediation (only if file exists)
medf=os.path.join(TAB,"mediation_WGD_6p_HLALOH.tsv")
if os.path.exists(medf):
    m6=load("mediation_WGD_6p_HLALOH.tsv")
    hdr=["Quantity","Estimate","95% CI low","95% CI high"]
    rows=[hdr]+[[r[0],r[1],r[2],r[3]] for r in m6[1:]]
    mktable(doc,rows,
      "Supplementary Table 6. Causal mediation of the whole-genome-doubling effect on HLA LOH through chromosome-6p loss.",
      "Nonparametric bootstrap (500 patient-level resamples) decomposition. ACME, average causal "
      "mediation (indirect) effect via 6p loss; ADE, average direct effect; proportion mediated = "
      "ACME/total effect. Estimates are differences in the probability of HLA LOH.",
      fs=8)
    print("Added ST5 + ST6")
else:
    print("Added ST5 (ST6 mediation pending)")

# ST7: germline zygosity x evaluability
z7=os.path.join(TAB,"SuppTable7_zygosity_evaluability.tsv")
if os.path.exists(z7):
    d7=load("SuppTable7_zygosity_evaluability.tsv")
    hdr=["Locus","Heterozygous N","Heteroz. evaluable %","Homozygous N","Homoz. evaluable %"]
    rows=[hdr]+[r for r in d7[1:]]
    mktable(doc,rows,
      "Supplementary Table 7. Germline HLA zygosity and locus-level LOH evaluability.",
      "Among genotyped patients, germline-homozygous class I loci are structurally non-evaluable "
      "for loss of heterozygosity (0% yield a call), because LOH is defined by loss of one of two "
      "distinct germline alleles. Heterozygous loci are evaluable in ~60% of tumors.",
      fs=8)
    print("Added ST7")

doc.save(os.path.join(OUTd,"Supplementary_Tables.docx"))
print("Saved Supplementary_Tables.docx with", len(doc.tables), "tables")
