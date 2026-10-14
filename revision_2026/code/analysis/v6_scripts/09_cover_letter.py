#!/usr/bin/env python3
"""Convert Cover Letter markdown to DOCX."""
import os, re
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
CL="C:/Users/Administrator/科研文件/HLA_LOH02/V6/09_cover_letter"
NAVY=RGBColor(0x00,0x33,0x66)
t=open(os.path.join(CL,"Cover_Letter.md"),encoding="utf-8").read()
doc=Document(); s=doc.styles["Normal"]; s.font.name="Times New Roman"; s.font.size=Pt(11)
for para in t.split("\n"):
    p=para.strip()
    if not p:
        doc.add_paragraph(); continue
    # bold markdown **...**
    para_obj=doc.add_paragraph(); para_obj.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
    parts=re.split(r'(\*\*.*?\*\*)', p)
    for part in parts:
        if part.startswith("**") and part.endswith("**"):
            r=para_obj.add_run(part[2:-2]); r.bold=True; r.font.color.rgb=NAVY
        else:
            para_obj.add_run(part)
doc.save(os.path.join(CL,"Cover_Letter.docx"))
print("Saved Cover_Letter.docx")
