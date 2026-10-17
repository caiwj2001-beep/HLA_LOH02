#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""verify_supp_v4.py — 交付后自检：补充 DOCX + 补充 PDF。

检查项
  [1] PDF 页数、每个锚点（S1..S14 + Figure S1/S2）落在哪一页
  [2] PDF 第 1 页目录表里写的页码 == 该条目实际所在页（两遍法的真伪判据）
  [3] DOCX 内 "0.000"/"0.0000"/"0.000-..." 的计数（应为 0）
  [4] DOCX 内 S11 的 clonal 行、S14 是否存在/行数/关键措辞
输出：supp_pdf_verify.json + 屏显摘要
"""
import os, re, json, sys
from docx import Document
from pypdf import PdfReader

V4 = r"C:\Users\Administrator\科研文件\HLA_LOH02\iscience投稿版本\V4"
DOCX = os.path.join(V4, "03_tables", "Supplementary_Tables.docx")
PDF  = os.path.join(V4, "04_supplement", "Supplemental_Information.pdf")
OUT  = os.path.join(V4, "QC", "P2", "supp_pdf_verify.json")

ZERO = re.compile(r"(?<![\d.])0\.0{3,}(?![\d])")
ITEM = re.compile(r"Supplementary (Figure|Table) S(\d+)")

def anchors_from_pdf(rd):
    pages = [(rd.pages[i].extract_text() or "") for i in range(len(rd.pages))]
    anchors = {}
    for kind in ("Figure", "Table"):
        for n in (range(1, 3) if kind == "Figure" else range(1, 15)):
            key = f"Supplementary {kind} S{n}"
            tgt = key + "."
            for i, t in enumerate(pages):
                if tgt in t:
                    anchors[key] = i + 1
                    break
    return anchors, pages

def toc_from_pdf(pages):
    t = pages[0]
    # the table of contents ends where the footnote legend begins; without this cut the
    # final row swallows the legend ("... map to S8 to S13 ... S14 ...") and the parser
    # would read the 14 of "S14" as that row's page number.
    cut = t.find("Supplementary tables are numbered continuously")
    if cut > 0:
        t = t[:cut]
    ms = list(ITEM.finditer(t))
    toc = {}
    for i, m in enumerate(ms):
        key = f"Supplementary {m.group(1)} S{int(m.group(2))}"
        seg = t[m.end(): ms[i + 1].start() if i + 1 < len(ms) else len(t)]
        nums = re.findall(r"(?<![\d.])(\d+)(?![\d.])", seg)
        toc[key] = int(nums[-1]) if nums else None
    return toc

def docx_zero_scan(doc):
    hits = []
    for ti, tb in enumerate(doc.tables):
        for ri, row in enumerate(tb.rows):
            for c in row.cells:
                if ZERO.search(c.text):
                    hits.append(("table", ti, ri, c.text.strip()[:80]))
    for p in doc.paragraphs:
        if ZERO.search(p.text):
            hits.append(("para", None, None, p.text.strip()[:80]))
    return hits

def find_table(doc, needle):
    for tb in doc.tables:
        if tb.rows and needle in tb.rows[0].cells[0].text:
            return tb
        for row in tb.rows:
            if needle in row.cells[0].text:
                return tb
    return None

def main():
    rd = PdfReader(PDF)
    anchors, pages = anchors_from_pdf(rd)
    toc = toc_from_pdf(pages)
    doc = Document(DOCX)
    zero = docx_zero_scan(doc)

    # S11: the clonal row
    s11_row, s11_hdr = None, None
    for tb in doc.tables:
        hdr = [c.text.strip() for c in tb.rows[0].cells]
        if hdr[:3] == ["Kind", "Model / definition", "N"]:
            s11_hdr = hdr
            for row in tb.rows[1:]:
                if "clonal" in row.cells[1].text.lower():
                    s11_row = [c.text.strip() for c in row.cells]

    # S14: find heading paragraph + its table
    s14_heading = None
    for p in doc.paragraphs:
        if p.text.strip().startswith("Supplementary Table S14."):
            s14_heading = p.text.strip()
    s14_rows, s14_note = 0, ""
    for i, p in enumerate(doc.paragraphs):
        if s14_heading and p.text.strip() == s14_heading:
            alt = doc.paragraphs[i + 1].text.strip()
            s14_note = alt
            break
    # the S14 table = last table whose header starts with 'Model'
    for tb in doc.tables:
        hdr = [c.text.strip() for c in tb.rows[0].cells]
        if hdr[:3] == ["Model", "Covariate", "Df"]:
            s14_rows = len(tb.rows) - 1
            s14_first = [c.text.strip() for c in tb.rows[1].cells]
            s14_last = [c.text.strip() for c in tb.rows[-1].cells]

    mismatches = {k: (toc.get(k), anchors.get(k)) for k in anchors
                  if toc.get(k) != anchors.get(k)}
    res = {
        "pdf": PDF, "docx": DOCX,
        "pdf_pages": len(rd.pages),
        "n_anchors": len(anchors),
        "anchors": anchors,
        "toc_written": toc,
        "toc_vs_actual_mismatch": mismatches,
        "toc_all_match": not mismatches,
        "docx_zero_hits": zero,
        "docx_zero_count": len(zero),
        "s11_header": s11_hdr, "s11_clonal_row": s11_row,
        "s14_heading": s14_heading, "s14_note_len": len(s14_note),
        "s14_rows": s14_rows,
        "s14_note_has_no_protect": ("must NOT be interpreted as protection" in s14_note)
                                   or ("MUST NOT be interpreted" in s14_note),
        "s14_note_has_source": "vif_diagnostics.tsv" in s14_note,
    }
    with open(OUT, "w", encoding="utf-8") as f:
        json.dump(res, f, indent=1, ensure_ascii=False)

    print(f"[1] PDF pages = {res['pdf_pages']} ; anchors = {len(anchors)}/16")
    for k, v in anchors.items():
        print(f"      {k:28s} -> page {v}")
    print(f"[2] TOC (page 1) vs actual: {'ALL MATCH' if res['toc_all_match'] else mismatches}")
    for k in anchors:
        print(f"      {k:28s} toc={toc.get(k)}  actual={anchors.get(k)}")
    print(f"[3] DOCX zero-P hits = {len(zero)}")
    for h in zero:
        print("      ", h)
    print(f"[4] S11 clonal row = {s11_row}")
    print(f"    S14 heading     = {s14_heading}")
    print(f"    S14 rows        = {s14_rows} ; source cited = {res['s14_note_has_source']} ; "
          f"no-protection clause = {res['s14_note_has_no_protect']}")
    ok = (res['pdf_pages'] >= 20 and len(anchors) == 16 and res['toc_all_match']
          and len(zero) == 0 and s14_rows == 23 and s11_row is not None)
    print(f"\nVERDICT: {'PASS' if ok else 'FAIL'}")
    return 0 if ok else 1

if __name__ == "__main__":
    sys.exit(main())
